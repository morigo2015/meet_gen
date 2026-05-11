from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook


@pytest.fixture
def make_workbook(tmp_path):
    def _make(headers, rows, sheet_name="Sheet1") -> Path:
        path = tmp_path / "data.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_name
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        return path

    return _make


@pytest.fixture
def make_template(tmp_path):
    def _make(paragraphs) -> Path:
        path = tmp_path / "template.docx"
        document = Document()
        for text in paragraphs:
            document.add_paragraph(text)
        document.save(path)
        return path

    return _make
